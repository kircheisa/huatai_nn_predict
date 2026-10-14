# -*- coding: utf-8 -*-
"""
导出 图文版 Word 报告：模型迭代报告_BiGRU.md → 模型迭代报告_BiGRU.docx
======================================================================
- 解析 markdown（标题/表格/代码块/列表/粗体/行内代码/引用/链接），映射为 Word 样式
- 第八章"配套可视化图表"自动嵌入 figs/ 下 9 张图（PNG 透明版）+ 图注
- 中文正文字体：微软雅黑；代码：Consolas；表格：Table Grid
- 运行: python 04_导出Word报告.py   （依赖: pip install python-docx）
"""

import re
import sys
from pathlib import Path

# 沙箱环境: python-docx 安装在项目内 _libs 目录
sys.path.insert(0, str(Path(__file__).parent / "_libs"))
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = Path(r"f:\43_华泰证券预处理\43_华泰证券预处理\model_nn")
MD = BASE / "模型迭代报告_BiGRU.md"
OUT = BASE / "模型迭代报告_BiGRU.docx"
FIGS = BASE / "figs"

# 第八章插入的图片（文件名前缀 → 图注）
PICTURES = [
    ("fig01_原始数据样本图",     "图1  原始数据样本概览：目标序列 / 分布 / 特征空间"),
    ("fig02_预处理过程可视化",   "图2  预处理流程：10天标签生成 / 28-fold滚动切分 / 死区加权"),
    ("fig03_模型结构示意图",     "图3  BiGRU + Inception + Self-Attention 模型结构与训练协议"),
    ("fig04_预测对比_时间序列",  "图4  预测 vs 真实：时间序列对比（概率、阈值与正误标记）"),
    ("fig05_散点拟合分布",       "图5  预测 vs 真实：散点拟合与概率校准（含 MSE/R²/LogLoss/AUC）"),
    ("fig06_误差热力图",         "图6  误差热力图：混淆矩阵与分年方向准确率"),
    ("fig07_指标雷达图",         "图7  模型评估指标雷达图（五维对比）"),
    ("fig08_参数性能对比柱状图", "图8  不同参数下的性能对比（horizon / seq_len / 训练窗口）"),
    ("fig09_误差分布箱线图",     "图9  误差分布箱线图：fold级稳定性与分年概率误差"),
]

MONO = "Consolas"
CN = "微软雅黑"


def set_cn(run, name=CN):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_inline(par, text: str, mono=False, size=10.5):
    """解析行内格式: **粗体** `代码` [文本](链接)"""
    for tok in re.split(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))", text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = par.add_run(tok[2:-2]); r.bold = True
            set_cn(r); r.font.size = Pt(size)
        elif tok.startswith("`") and tok.endswith("`"):
            r = par.add_run(tok[1:-1]); r.font.name = MONO; r.font.size = Pt(size - 1)
            r.font.color.rgb = RGBColor(0x8B, 0x2F, 0x2F)
        elif tok.startswith("[") and "](" in tok:
            label = tok[1:tok.index("]")]
            r = par.add_run(label); set_cn(r); r.font.size = Pt(size)
            r.font.color.rgb = RGBColor(0x1F, 0x5C, 0x99)
        else:
            r = par.add_run(tok); set_cn(r); r.font.size = Pt(size)
            if mono:
                r.font.name = MONO


def add_table(doc: Document, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{3,}:?", c or "---") for c in r)]
    ncol = max(len(r) for r in cells)
    t = doc.add_table(rows=len(cells), cols=ncol)
    t.style = "Table Grid"
    for i, row in enumerate(cells):
        for j in range(ncol):
            txt = row[j] if j < len(row) else ""
            cell = t.cell(i, j)
            cell.paragraphs[0].text = ""
            add_inline(cell.paragraphs[0], txt, size=9.5)
            if i == 0:
                for r in cell.paragraphs[0].runs:
                    r.bold = True
    doc.add_paragraph()


def add_code(doc: Document, lines):
    for ln in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(ln if ln else " ")
        r.font.name = MONO; r.font.size = Pt(9)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), CN)
        r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    doc.add_paragraph()


def add_pictures(doc: Document):
    for stem, caption in PICTURES:
        f = FIGS / f"{stem}.png"
        if not f.exists():
            continue
        doc.add_picture(str(f), width=Inches(6.3))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(caption)
        set_cn(r); r.font.size = Pt(9.5); r.bold = True
        r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)


def main():
    lines = MD.read_text(encoding="utf-8").splitlines()
    doc = Document()
    # 全局样式
    st = doc.styles["Normal"]
    st.font.name = "Calibri"; st.font.size = Pt(10.5)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), CN)
    for h, sz in [("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 11.5)]:
        s = doc.styles[h]
        s.font.name = "Calibri"; s.font.size = Pt(sz); s.font.bold = True
        s.font.color.rgb = RGBColor(0x1F, 0x3B, 0x5C)
        rpr = s._element.get_or_add_rPr().get_or_add_rFonts()
        rpr.set(qn("w:eastAsia"), CN)

    i, in_sec8, sec8_done = 0, False, False
    in_code, code_buf = False, []
    while i < len(lines):
        ln = lines[i]

        if ln.strip().startswith("```"):
            if in_code:
                add_code(doc, code_buf); code_buf = []
            in_code = not in_code
            i += 1; continue
        if in_code:
            code_buf.append(ln); i += 1; continue

        s = ln.strip()
        if s.startswith("|"):
            tbl = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                tbl.append(lines[i]); i += 1
            add_table(doc, tbl); continue
        if not s:
            i += 1; continue
        if s == "---":
            if in_sec8 and not sec8_done:      # 第八章结束处插入全部图片
                doc.add_page_break()
                h = doc.add_heading("图表合集（figs/，可二次编辑的 SVG 矢量版同步输出）", 2)
                add_pictures(doc)
                sec8_done = True; in_sec8 = False
            i += 1; continue

        if s.startswith("# "):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(s[2:]); r.bold = True; set_cn(r)
            r.font.size = Pt(19); r.font.color.rgb = RGBColor(0x14, 0x2C, 0x4A)
            doc.add_paragraph()
        elif s.startswith("## "):
            doc.add_heading(s[3:], 1)
            in_sec8 = s[3:].startswith("八")
        elif s.startswith("### "):
            doc.add_heading(s[4:], 2)
        elif s.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_inline(p, s[2:], size=9.5)
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x5A, 0x5A, 0x5A)
        elif re.match(r"^\d+\.\s", s) or s.startswith("- "):
            p = doc.add_paragraph(style="List Number" if s[0].isdigit() else "List Bullet")
            add_inline(p, re.sub(r"^(\d+\.|-)\s+", "", s))
        elif s.startswith("*") and s.endswith("*"):
            p = doc.add_paragraph()
            r = p.add_run(s.strip("*")); set_cn(r)
            r.italic = True; r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
        else:
            p = doc.add_paragraph()
            add_inline(p, s)
        i += 1

    from datetime import datetime
    out_path = OUT
    try:
        doc.save(out_path)
    except PermissionError:
        out_path = OUT.with_name(f"模型迭代报告_BiGRU_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx")
        doc.save(out_path)
    print(f"已生成: {out_path}", flush=True)
    print(f"插图: {sum(1 for s, _ in PICTURES if (FIGS / f'{s}.png').exists())} 张", flush=True)


if __name__ == "__main__":
    main()
