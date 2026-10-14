# -*- coding: utf-8 -*-
"""从已保存的实验产物生成15页国债方向预测汇报，不重新训练模型。"""
from io import BytesIO
import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import PercentFormatter
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from sklearn.metrics import accuracy_score, f1_score, roc_auc_score

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
OUT = HERE / "out"
DEST = ROOT / "训练效果总结"
PPT = DEST / "国债收益率方向预测_模型演进与优化汇报_20261014.pptx"
FONT = "Microsoft YaHei"
NAVY, TEAL, GOLD, RED = "10243A", "087F8C", "C38B22", "C9474D"
INK, MUTED, LIGHT, WHITE = "26374A", "64758B", "F3F6FA", "FFFFFF"
plt.rcParams.update({
    "font.sans-serif": [FONT], "axes.unicode_minus": False,
    "font.size": 12, "axes.labelcolor": "#" + INK,
    "text.color": "#" + INK, "xtick.color": "#" + MUTED,
    "ytick.color": "#" + MUTED, "axes.spines.top": False,
    "axes.spines.right": False, "axes.edgecolor": "#D5DDE7",
    "figure.facecolor": "white", "axes.facecolor": "white",
})


def read_json(name):
    return json.loads((OUT / name).read_text(encoding="utf-8"))


def metrics(y, p, pred=None):
    y, p = np.asarray(y), np.asarray(p)
    pred = (p >= .5).astype(int) if pred is None else np.asarray(pred)
    return {"n": len(y), "acc": accuracy_score(y, pred),
            "f1": f1_score(y, pred, average="macro", zero_division=0),
            "auc": roc_auc_score(y, p) if len(np.unique(y)) > 1 else np.nan}


def pct(v):
    return f"{v:.2%}"


def load_inputs():
    audit = read_json("OPT9_validation.json")
    rate = pd.read_csv(ROOT / "data_ori" / "2026_origin_rate.csv").iloc[:, :2]
    rate.columns = ["date", "rate"]
    rate["date"] = pd.to_datetime(rate["date"], errors="coerce")
    rate["rate"] = pd.to_numeric(rate["rate"], errors="coerce")
    rate = rate.dropna().sort_values("date").drop_duplicates("date").reset_index(drop=True)
    rate["delta_bp"] = (rate["rate"].shift(-10) - rate["rate"]) * 100
    rate["maturity"] = rate["date"].shift(-10)
    rates = rate.set_index("date")
    predictions = {}
    for tag in audit:
        p = pd.read_csv(OUT / f"{tag}_predictions.csv", parse_dates=["date"])
        p["delta_bp"] = p["date"].map(rates["delta_bp"])
        assert p["delta_bp"].notna().all()
        p["original_true"] = (p["delta_bp"] < 0).astype(int)
        p["correct"] = p["pred"] == p["original_true"]
        p["dev_mask"] = ((p["date"] <= "2024-12-31")
                         & (p["date"].map(rates["maturity"]) <= "2024-12-31"))
        p["year"] = p["date"].dt.year
        p["month"] = p["date"].dt.month
        predictions[tag] = p
    return audit, predictions, rate


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    if hasattr(shape, "line"):
        shape.line.fill.background()


def box(slide, x, y, w, h, color=WHITE, rounded=True):
    s = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    fill(s, color)
    if rounded:
        s.adjustments[0] = .08
    return s


def text(slide, value, x, y, w, h, size=18, color=INK, bold=False,
         align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.02)
    tf.margin_top = tf.margin_bottom = Inches(.015)
    for i, line in enumerate(str(value).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(6)
        r = p.add_run()
        r.text = line
        r.font.name, r.font.size, r.font.bold = FONT, Pt(size), bold
        r.font.color.rgb = RGBColor.from_string(color)
        ea = OxmlElement("a:ea")
        ea.set("typeface", FONT)
        r._r.get_or_add_rPr().append(ea)
    return shape


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.titles = []

    def slide(self, title, subtitle, source, notes=""):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        fill(slide.background, LIGHT)
        box(slide, 0, 0, 13.333, .12, TEAL, False)
        text(slide, "国债收益率方向预测  /  研究进展汇报", .55, .27, 10, .3, 10, MUTED)
        text(slide, title, .55, .78, 12.1, .62, 29, NAVY, True)
        text(slide, subtitle, .57, 1.48, 12.0, .5, 12, MUTED)
        text(slide, source, .57, 7.05, 11.7, .24, 8.5, MUTED)
        text(slide, f"{len(self.prs.slides):02d} / 15", 12.1, 7.01, .65, .3,
             10, MUTED, align=PP_ALIGN.RIGHT)
        slide.notes_slide.notes_text_frame.text = (
            f"{title}\n口径与来源：{source}\n{notes}")
        self.titles.append(title)
        return slide

    def chart(self, slide, fig, x, y, w, h):
        image = BytesIO()
        fig.savefig(image, format="png", dpi=190, bbox_inches="tight", facecolor="white")
        plt.close(fig)
        image.seek(0)
        slide.shapes.add_picture(image, Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def card(slide, x, y, w, h, label, value, detail, accent=TEAL):
    box(slide, x, y, w, h)
    box(slide, x, y, .055, h, accent, False)
    text(slide, label, x + .2, y + .2, w - .4, .4, 14, MUTED)
    text(slide, value, x + .2, y + .75, w - .4, .65, 31, accent, True)
    text(slide, detail, x + .2, y + 1.52, w - .4, h - 1.6, 14)


def bullet_panel(slide, x, y, w, h, heading, items, accent=TEAL, size=17):
    box(slide, x, y, w, h)
    compact = h < 2.4
    start = .58 if compact else .83
    text(slide, heading, x + .22, y + .12, w - .44, .38 if compact else .5,
         17 if compact else 20, accent, True)
    step = (h - start - .1) / max(len(items), 1)
    for i, item in enumerate(items):
        box(slide, x + .23, y + start + .12 + i * step, .055, .055, accent, False)
        text(slide, item, x + .43, y + start + i * step, w - .67, step - .025, size)


def table(slide, columns, rows, x, y, widths, row_h=.46, size=15,
          highlights=()):
    total = sum(widths)
    offset = .04 if row_h < .3 else .09
    text_h = row_h - .055 if row_h < .3 else row_h - .1
    box(slide, x, y, total, row_h, NAVY, False)
    for j, (name, width) in enumerate(zip(columns, widths)):
        text(slide, name, x + sum(widths[:j]) + .12, y + offset,
             width - .18, text_h, size, WHITE, True)
    for i, row in enumerate(rows):
        color = "E6F3F1" if i in highlights else (WHITE if i % 2 == 0 else "EAF0F6")
        yy = y + (i + 1) * row_h
        box(slide, x, yy, total, row_h, color, False)
        for j, (value, width) in enumerate(zip(row, widths)):
            text(slide, value, x + sum(widths[:j]) + .12, yy + offset,
                 width - .18, text_h, size, INK, i in highlights)


def main():
    audit, pred, rate = load_inputs()
    baseline_tag = "OPT9_featrelab_E"
    base = pred[baseline_tag]
    p25 = base[base.year == 2025].copy()
    oct25 = p25[p25.month == 10].copy()
    dev = audit[baseline_tag]["by_segment"]["dev"]["original_direction"]
    yr = audit[baseline_tag]["by_segment"]["2025"]["original_direction"]
    monthly = base.groupby(base.date.dt.to_period("M")).agg(
        n=("correct", "size"), acc=("correct", "mean"), probability=("proba_down", "mean"))
    worst = monthly[monthly.n >= 15].sort_values(["acc", "n"], ascending=[True, False]).head(6)
    deck = Deck()
    source = "来源：model_nn/out/OPT9_validation.json；预测明细按原始利率重新计算未来10日涨跌"

    # 01
    s = deck.slide("国债收益率方向预测", "从XGBoost到神经网络：特征、训练、优化与预测失误复盘",
                   "数据截至：现有滚动回测产物；汇报日期：2026年10月14日",
                   "本报告不是交易收益回测。当前基线是研究对照，不是已验证可部署的生产模型。")
    text(s, "模型演进与优化汇报", .65, 2.15, 11.9, .8, 39, NAVY, True)
    text(s, "历史探索发现潜力，因果复核界定边界，逐日诊断决定下一步", .68, 3.2,
         11.6, .65, 20, MUTED)
    card(s, .68, 4.26, 3.85, 2.2, "评估框架", "28折滚动", "1,509个开发期预测\n243个2025年预测")
    card(s, 4.74, 4.26, 3.85, 2.2, "当前E基线 · 2025原始方向", pct(yr["acc"]),
         "多数类60.08%；净增益2.88个百分点")
    card(s, 8.8, 4.26, 3.85, 2.2, "典型风险 · 2025年10月", "15 / 17日误判",
         "温和下行被漏判；事件特征未修复", RED)

    # 02
    s = deck.slide("结论先行：模型迁移已完成，稳定增益尚未确立",
                   "研究对照基线保留E特征＋down_ex2；不把2025单年最优直接作为部署依据", source)
    card(s, .65, 2.1, 3.87, 2.55, "历史探索的主要突破", "H5 → H10",
         "旧版NN全局AUC提升至0.6705\n不同预测周期，不作纯架构归因")
    card(s, 4.73, 2.1, 3.87, 2.55, "当前因果复核边界", pct(dev["acc"]),
         "开发期原始方向54.21%\n低于同口径多数类59.05%", RED)
    card(s, 8.81, 2.1, 3.87, 2.55, "新增事件方向标签", "2025 +1.23pt",
         "64.20%，但开发期下降3.05pt\n未通过预先固定的采用门槛", GOLD)
    box(s, .65, 5.03, 12.03, 1.55, NAVY)
    text(s, "下一阶段的目标", .9, 5.25, 2.7, .45, 19, WHITE, True)
    text(s, "从“某年准确率更高”转向“成熟标签、固定协议、稳定超基线”，\n先校准概率与完善时间可得性，再做同周期融合及交易成本评估。",
         3.4, 5.2, 8.95, 1.15, 19, WHITE)

    # 03
    s = deck.slide("特征体系：从缓存工程特征转向可追溯因果特征",
                   "收益率、宏观、资金、政策及机构行为按可用时点对齐；不把新闻描述直接当未来标签",
                   "来源：01_BiGRU训练_二分类_10天.py；05_基础特征对照训练.py；09_方向优化实验.py")
    table(s, ["阶段", "输入与构造", "特征规模", "状态 / 风险"], [
        ["历史XGB / NN", "STL、遗传规划、滚动工程特征；NN取XGB Top-30", "280 → 30", "全期筛选及缓存需复核"],
        ["D：因果基础", "日/周/月/季发布对齐＋收益率变化＋资金利差", "100列", "发布日＋1日；历史始于2010"],
        ["E：趋势扩展", "60/120日动量、水平z分数、回撤、长周期波动", "120列", "当前研究对照基线"],
        ["F：市场行为", "政策锚、FR007、机构净买入、供给和成交量", "144列", "验证未通过，非默认输入"],
        ["E＋客观事件", "七类事件、强度、20日密度、距事件天数", "130列", "不使用人工direction"],
        ["E＋方向事件", "再加利多/利空衰减和及净情绪", "133列", "回溯标注存在事后解释风险"],
    ], .65, 2.08, [1.65, 6.0, 1.4, 3.0], row_h=.61, size=13.5, highlights=(2,))
    text(s, "数值特征：过去252日中位数/IQR归一化，折内填补/筛列/标准化；附加缺失标记。\n事件特征：按完整交易日历构建，事件日之后首个交易日起生效；回溯精选日历≠实时完备新闻库。",
         .75, 6.45, 11.9, .52, 12.5, MUTED)

    # 04
    raw_x = read_json("RAW_B_xgb_S2_E40_metrics.json")
    raw_n = read_json("RAW_B_nn_S2_E40_metrics.json")
    s = deck.slide("XGB到神经网络：迁移不是自动提升",
                   "先分清预测周期和数据口径，再看同源对照；不能用H5与H10差异证明架构优越",
                   "来源：model_new/data/xgb_metrics_bin.txt、xgb_metrics_10d.txt；RAW_B_*_S2_E40_metrics.json")
    table(s, ["历史探索模型", "周期", "全局Acc", "全局AUC", "样本 / 比较边界"], [
        ["XGB缓存基线", "5日", "56.35%", "0.5822", "1,764；280特征"],
        ["XGB缓存H10", "10日", "57.74%", "0.6344", "1,827；29折、116特征"],
        ["旧版NN FINAL5", "10日", "59.13%", "0.6705", "1,764；28折、Top-30"],
    ], .65, 2.08, [3.55, 1.0, 1.5, 1.5, 4.5], row_h=.57, size=15)
    fig, ax = plt.subplots(figsize=(7.3, 2.3))
    names = ["同源H10 XGB", "同源H10 NN", "当前E研究基线"]
    vals = [raw_x["years"]["2025"]["acc"], raw_n["years"]["2025"]["acc"], yr["acc"]]
    ax.barh(names[::-1], vals[::-1], color=["#" + TEAL, "#6487A6", "#" + GOLD])
    ax.axvline(yr["majority_baseline"], ls="--", color="#" + RED, label="2025多数类60.08%")
    for i, v in enumerate(vals[::-1]):
        ax.text(v + .012, i, pct(v), va="center", fontsize=12)
    ax.set_xlim(0, .82); ax.xaxis.set_major_formatter(PercentFormatter(1))
    ax.set_xlabel("2025原始方向准确率"); ax.legend(loc="upper right", fontsize=9, frameon=False)
    fig.tight_layout()
    deck.chart(s, fig, .7, 4.62, 7.3, 2.02)
    bullet_panel(s, 8.25, 4.57, 4.35, 2.06, "同源H10对照的含义",
                 ["XGB与NN使用同源B特征、相同测试日历；不是同超参数对照。",
                  "全局Acc：XGB54.20%，NN52.21%；不能宣称全面替代。"], size=12.5)

    # 05
    s = deck.slide("神经网络架构：从复杂堆叠到小样本CompactGRU",
                   "历史版负责探索序列模式；当前版通过降容量与固定协议建立更清晰的对照",
                   "来源：01_BiGRU训练_二分类_10天.py；05_基础特征对照训练.py:CompactGRU")
    for xx, heading, steps, accent in [
        (.65, "历史版：BiGRU＋Inception＋Attention",
         ["40日 × Top-30输入", "双向GRU：hidden=64", "四分支卷积：k=3/5/7/9", "Self-Attention＋全局池化", "二分类logits；5-seed概率平均"], GOLD),
        (6.8, "当前版：CompactGRU（09实验）",
         ["20日 × 120列E特征＋缺失标记", "Linear＋LayerNorm＋GELU投影", "单层双向GRU：hidden=32", "均值池化＋末端状态拼接", "Dropout＋二分类；3-seed概率平均"], TEAL),
    ]:
        box(s, xx, 2.06, 5.82, 4.56)
        text(s, heading, xx + .18, 2.25, 5.46, .45, 18, accent, True)
        for i, step in enumerate(steps):
            yy = 2.99 + i * .65
            box(s, xx + .3, yy, 5.22, .47, LIGHT)
            text(s, step, xx + .43, yy + .065, 4.95, .37, 16,
                 accent if i == 0 else INK, i == 0)
    text(s, "当前实验没有Inception或Attention模块；不能把旧架构图当作当前训练结构。",
         .8, 6.68, 11.8, .3, 13, RED, True)

    # 06
    s = deck.slide("训练与预测流程：逐折只使用已成熟标签",
                   "预测日t输出未来10个交易日方向；标签到t＋10才落地；训练、验证、测试严格按时间切分",
                   "来源：09_方向优化实验.py:run_task；01_BiGRU训练_二分类_10天.py:train_one_fold")
    flow = [("构建样本", "过去20日序列"), ("成熟标签过滤", "成熟日早于测试首日"),
            ("训练 / 验证", "时间顺序＋20日purge"), ("3-seed训练", "验证AUC早停"),
            ("滚动测试", "保存概率与日期")]
    for i, (heading, body) in enumerate(flow):
        xx = .65 + i * 2.45
        box(s, xx, 2.22, 2.2, 1.4, NAVY if i == 3 else WHITE)
        text(s, heading, xx + .15, 2.47, 1.95, .4, 18, WHITE if i == 3 else TEAL, True)
        text(s, body, xx + .15, 3.01, 1.95, .48, 13, WHITE if i == 3 else MUTED)
        if i < 4:
            text(s, "→", xx + 2.19, 2.72, .28, .45, 21, MUTED)
    bullet_panel(s, .65, 4.02, 5.86, 2.72, "训练设置（当前E与事件候选）",
                 ["扩展历史窗口；lr=5e-4；dropout=0.3。",
                  "25轮上限；patience=8；AdamW、余弦退火、梯度裁剪。",
                  "无类别加权、无标签平滑；固定阈值0.5。"], size=15)
    bullet_panel(s, 6.76, 4.02, 5.86, 2.72, "评估与采用门槛",
                 ["开发期：预测及标签成熟日均≤2024-12-31。",
                  "采用：原始方向超多数类、配对增益区间为正、多数seed改善。",
                  "2025已多轮查看，仅作诊断；新时期盲测仍需补齐。"], size=15)

    # 07
    s = deck.slide("历史优化：预测周期比单纯增大模型更重要",
                   "旧版缓存特征探索结果；这些数值不与最新因果实验直接拼接为同口径排名",
                   "来源：LS2、SEQ60_ACC3、H10_SEQ40、H10_S40_FINAL5指标与预测产物；历史迭代报告")
    files = ["LS2", "SEQ60_ACC3", "H10_SEQ40", "H10_S40_FINAL5"]
    hist_rows = []
    for name in files:
        p = pd.read_csv(OUT / f"{name}_nn_test_predictions.csv")
        m = metrics(p.y_true, p.proba_down, p.pred)
        hist_rows.append((name, m))
    fig, ax = plt.subplots(figsize=(8.4, 3.6))
    names = ["NN H5\nLS2", "NN H5\nSEQ60", "NN H10\nSEQ40", "NN H10\nFINAL5"]
    x = np.arange(4)
    ax.bar(x - .17, [r[1]["acc"] for r in hist_rows], .33, label="准确率", color="#" + TEAL)
    ax.bar(x + .17, [r[1]["auc"] for r in hist_rows], .33, label="AUC", color="#6487A6")
    ax.set_xticks(x, names); ax.set_ylim(.45, .76); ax.legend(frameon=False, ncol=2)
    for i, (_, m) in enumerate(hist_rows):
        ax.text(i - .17, m["acc"] + .008, f"{m['acc']:.3f}", ha="center", fontsize=11)
        ax.text(i + .17, m["auc"] + .008, f"{m['auc']:.3f}", ha="center", fontsize=11)
    ax.set_title("缓存特征探索：全局Acc与AUC", loc="left", fontsize=15)
    fig.tight_layout(); deck.chart(s, fig, .65, 2.15, 8.0, 3.98)
    bullet_panel(s, 8.88, 2.15, 3.75, 3.98, "得到的经验",
                 ["Top-30优于堆砌全部特征，控制模型容量。",
                  "H10＋40日序列出现更强排序能力。",
                  "多seed减方差，但不保证准确率单调提升。"], size=15)
    text(s, "阈值教训：旧模型W=250滚动校准的2025准确率为48.15%，不是早期误算的78.60%。\n历史th=0.39的全局汇总包含开发样本选阈值，不能视为全局完全样本外校准收益。",
         .75, 6.31, 11.85, .55, 13, RED)

    # 08
    s = deck.slide("当前优化结果：标签与趋势有局部改善，泛化仍不足",
                   "同一原始方向口径；重定义标签的高原生准确率不能解释为方向预测准确率", source)
    tags = ["base_D", "relabel_D", "twostage_D", "features_E", "featrelab_E"]
    rows = []
    for name in tags:
        d = audit["OPT9_" + name]["by_segment"]
        rows.append([name, pct(d["dev"]["original_direction"]["acc"]),
                     pct(d["2025"]["original_direction"]["acc"]),
                     pct(d["2025"]["native"]["acc"])])
    table(s, ["实验", "开发期原始方向", "2025原始方向", "2025原生标签"], rows,
          .65, 2.1, [3.9, 2.7, 2.7, 2.7], .6, 16, highlights=(4,))
    box(s, .65, 6.0, 5.9, .83)
    text(s, "E基线是研究对照，不是生产模型", .86, 6.1, 5.45, .32, 16, RED, True)
    text(s, "开发期54.21%仍低于多数类59.05%。", .86, 6.49, 5.45, .26, 12)
    text(s, "两阶段未通过：幅度识别开发期AUC约0.466，过滤不能稳定增益。\n原生标签down_ex2：未来10日下跌超过2bp；原始方向：未来10日变化<0。",
         6.9, 6.02, 5.65, .82, 12.5, MUTED)

    # 09
    s = deck.slide("训练侧与市场特征优化：单年最优不等于可采用",
                   "以E基线为对照；短窗口、近期加权、状态相似度和F特征均未提供稳定开发期收益", source)
    names = ["featrelab_E", "featrelab_E_W1260", "featrelab_E_W756", "featrelab_E_W504",
             "featrelab_E_R252", "featrelab_E_S20", "featF_relab", "featF_down"]
    labels = ["E基线", "窗口1260", "窗口756", "窗口504", "近期加权252", "状态相似20", "F＋重定义标签", "F＋原始方向"]
    values = np.array([[audit["OPT9_" + t]["by_segment"][seg]["original_direction"]["acc"]
                        for seg in ("dev", "2025")] for t in names])
    fig, ax = plt.subplots(figsize=(8.4, 4.8))
    im = ax.imshow(values, cmap="RdYlGn", vmin=.4, vmax=.68, aspect="auto")
    ax.set_xticks([0, 1], ["开发期原始方向", "2025原始方向"])
    ax.set_yticks(np.arange(len(labels)), labels)
    for i in range(len(labels)):
        for j in range(2):
            ax.text(j, i, pct(values[i, j]), ha="center", va="center", fontsize=13)
    fig.colorbar(im, ax=ax, shrink=.75).ax.yaxis.set_major_formatter(PercentFormatter(1))
    fig.tight_layout(); deck.chart(s, fig, .7, 2.04, 7.64, 4.61)
    bullet_panel(s, 8.65, 2.1, 3.97, 4.47, "为什么不采用",
                 ["短窗：开发期45.79%–51.69%，丢失长历史信息的可能性需验证。",
                  "R252：2025达65.84%，但开发期仅50.10%，不能凭单年选择。",
                  "F：机构/资金/供给特征未修复10月；不等于已证明无预测信息。"], size=14)

    # 10
    s = deck.slide("政策事件消融：排序改善，准确率收益没有稳定复现",
                   "72个精选回溯事件；半衰期10个交易日；次交易日起可用；两个候选均未通过采用门槛",
                   "来源：02_文本特征构建.py；OPT9_event_*；文本标签_政策事件日历.csv")
    event_tags = [baseline_tag, "OPT9_event_objective_E", "OPT9_event_direction_E"]
    rows = []
    for name, tag in zip(["E基线", "E＋客观事件（10列）", "E＋方向事件（13列）"], event_tags):
        r = audit[tag]["by_segment"]
        rows.append([name, pct(r["dev"]["original_direction"]["acc"]),
                     pct(r["2025"]["original_direction"]["acc"]),
                     f"{r['2025']['original_direction']['auc']:.4f}", "11.76%"])
    table(s, ["模型", "开发期Acc", "2025 Acc", "2025 AUC", "2025年10月"], rows,
          .65, 2.11, [4.1, 2.0, 2.0, 2.0, 2.0], .64, 16)
    bullet_panel(s, .65, 5.06, 5.9, 1.68, "方向标签组：不以2025选型",
                 ["2025净增益1.23pt，95%分块区间−11.11至＋13.59pt。",
                  "开发期下降3.05pt；仅1/3个seed在开发期改善。"], accent=RED, size=14)
    bullet_panel(s, 6.8, 5.06, 5.82, 1.68, "避免把关联当机制证明",
                 ["2025事件后20日准确率：基线61.27%，方向组56.34%。",
                  "10月20日新增事件仅能从10月21日起影响预测。"], size=14)

    # 11
    s = deck.slide("结果分布：准确率主要来自非下行日，漏判风险集中",
                   "当前E基线 · 2025年243日；预测概率p对应“下跌超过2bp”，不是纯方向概率", source)
    fig, axes = plt.subplots(1, 2, figsize=(11.5, 3.5))
    ax = axes[0]
    ax.hist(p25.loc[p25.original_true == 0, "proba_down"], bins=np.linspace(0, 1, 16),
            alpha=.75, color="#" + TEAL, label="实际上行/持平")
    ax.hist(p25.loc[p25.original_true == 1, "proba_down"], bins=np.linspace(0, 1, 16),
            alpha=.65, color="#" + RED, label="实际下行")
    ax.axvline(.5, color="#" + NAVY, ls="--")
    ax.set_xlabel("模型输出p（下跌超过2bp）"); ax.set_ylabel("预测日数量")
    ax.set_title("概率分布与0.5判定线", loc="left"); ax.legend(fontsize=10)
    ax = axes[1]
    cm = pd.crosstab(p25.original_true, p25.pred).reindex(index=[0, 1], columns=[0, 1], fill_value=0).to_numpy()
    ax.imshow(cm, cmap="Blues")
    ax.set_xticks([0, 1], ["预测非下行", "预测下行"])
    ax.set_yticks([0, 1], ["实际非下行", "实际下行"])
    for i in range(2):
        for j in range(2):
            ax.text(j, i, str(cm[i, j]), ha="center", va="center", fontsize=28,
                    color="white" if cm[i, j] > 90 else "#" + INK)
    ax.set_title("折叠为原始方向后的混淆矩阵", loc="left")
    fig.tight_layout(); deck.chart(s, fig, .7, 2.07, 11.85, 3.55)
    box(s, .7, 5.87, 3.77, .88)
    text(s, "预测下行 / 实际下行", .89, 5.98, 3.4, .27, 13, MUTED)
    text(s, "29日 / 97日", .89, 6.3, 3.4, .38, 23, RED, True)
    text(s, "下行召回率：18/97＝18.56%；漏判79日。\n总体Acc62.96%，但MacroF1仅0.5179、原始方向AUC0.5086。",
         4.85, 6.03, 7.65, .68, 16, INK)

    # 12
    s = deck.slide("哪些时期预测不好：按月看存在成片失误",
                   "E基线原始方向准确率；历史最差月份按n≥15筛选，非事前可执行的过滤规则", source)
    heat = base[base.year <= 2025].pivot_table(
        index="year", columns="month", values="correct", aggfunc="mean").reindex(columns=range(1, 13))
    fig, ax = plt.subplots(figsize=(10.5, 3.0))
    im = ax.imshow(heat.to_numpy(dtype=float), cmap="RdYlGn", vmin=0, vmax=1, aspect="auto")
    ax.set_xticks(np.arange(12), [f"{i}月" for i in range(1, 13)])
    ax.set_yticks(np.arange(len(heat)), heat.index)
    for i in range(len(heat)):
        for j in range(12):
            v = heat.iloc[i, j]
            if pd.notna(v):
                ax.text(j, i, f"{v:.0%}", ha="center", va="center", fontsize=8.5)
    fig.colorbar(im, ax=ax, shrink=.85).ax.yaxis.set_major_formatter(PercentFormatter(1))
    fig.tight_layout(); deck.chart(s, fig, .66, 2.05, 11.94, 2.88)
    rows = [[str(period), int(row.n), pct(row.acc), f"{row.probability:.3f}"]
            for period, row in worst.iterrows()]
    table(s, ["最差月份", "样本数", "原始方向Acc", "平均模型p"], rows,
          .72, 5.1, [1.7, 1.05, 1.8, 1.65], row_h=.24, size=9.5)
    bullet_panel(s, 7.24, 5.09, 5.38, 1.7, "诊断而非因果定论",
                 ["2020-06、2023-04、2024-12等出现整月失误。",
                  "预测日共享10日标签区间，错误聚集并非独立重复证据。",
                  "体制切换、校准漂移、标签目标错配均需分开验证。"], size=11.5)

    # 13
    s = deck.slide("2025年10月复盘：温和下行被系统性漏判",
                   "17个预测日中15日错误；模型全部预测非下行；不能由单段失误认定历史无类似样本", source)
    fig, ax = plt.subplots(figsize=(7.9, 3.7))
    xx = np.arange(len(oct25))
    ax.bar(xx, oct25.delta_bp, color=np.where(oct25.delta_bp < 0, "#" + RED, "#" + TEAL), alpha=.65)
    ax.axhline(0, color="#8796AA", lw=.8); ax.set_ylabel("未来10日收益率变化（bp）")
    ax.set_xticks(xx[::2], oct25.date.dt.strftime("%m-%d").iloc[::2], rotation=35)
    a2 = ax.twinx()
    a2.plot(xx, oct25.proba_down, color="#" + NAVY, marker="o", lw=2, label="模型p")
    a2.axhline(.5, ls="--", color="#" + GOLD); a2.set_ylim(0, .72)
    a2.set_ylabel("模型p（下跌超过2bp）")
    ax.set_title("实际变化为负，概率仍持续低于0.5", loc="left", fontsize=13)
    fig.tight_layout(); deck.chart(s, fig, .67, 2.07, 7.8, 3.85)
    show_dates = ["2025-10-16", "2025-10-20", "2025-10-21", "2025-10-22", "2025-10-24"]
    selected = oct25.set_index(oct25.date.dt.strftime("%Y-%m-%d")).loc[show_dates]
    rows = [[d[-5:], f"{r.proba_down:.3f}", f"{r.delta_bp:+.2f}"]
            for d, r in selected.iterrows()]
    table(s, ["预测日", "模型p", "实际变化bp"], rows, 8.77, 2.21, [1.15, 1.15, 1.6],
          .45, 13)
    text(s, "10-09至10-29共15个连续\n预测日漏判；10-30、10-31正确。",
         8.85, 5.1, 3.78, .72, 14, RED, True)
    text(s, "目标错配只解释部分失误：−0.07至−1.99bp小跌在down_ex2中属于负类；\n但10-16、10-20等超过2bp的下行仍被漏判。F特征、客观事件和方向事件均未修复。",
         .76, 6.19, 11.75, .62, 14, INK)

    # 14
    s = deck.slide("不确定性与证据边界：不靠单年或单seed拍板",
                   "连续20日分块自助区间用于缓解重叠标签相关性；区间包含零，不等于稳定增益", source)
    ci_tags = ["OPT9_event_objective_E", "OPT9_event_direction_E"]
    fig, axes = plt.subplots(1, 2, figsize=(10.7, 3.15))
    for ax, segment, title in zip(axes, ["dev", "2025"], ["开发期：相对E基线增益", "2025：相对E基线增益"]):
        for i, tag in enumerate(ci_tags):
            r = audit[tag]["by_segment"][segment]
            lo, hi = r["original_vs_E_ci_block20"]
            v = r["original_vs_E"]
            ax.plot([lo * 100, hi * 100], [i, i], lw=3, color="#" + (TEAL if i == 0 else GOLD))
            ax.scatter(v * 100, i, s=65, color="#" + NAVY, zorder=3)
        ax.axvline(0, ls="--", color="#" + RED); ax.set_yticks([0, 1], ["客观事件", "方向事件"])
        ax.set_ylim(-.5, 1.5); ax.set_xlabel("准确率差（百分点）"); ax.set_title(title, loc="left", fontsize=13)
    fig.tight_layout(); deck.chart(s, fig, .72, 2.15, 11.8, 3.0)
    bullet_panel(s, .72, 5.49, 5.73, 1.36, "已经确认",
                 ["18项测试通过；日期、折边界、3-seed与事件指纹一致。",
                  "原生标签、原始方向、同子集多数类分开报告。"], size=12)
    bullet_panel(s, 6.74, 5.49, 5.87, 1.36, "仍未确认",
                 ["2025已被反复观察；后续需要真正未调参的新时期。",
                  "精选回溯事件、旧缓存筛选、宏观修订值均需可得性审计。"], accent=RED, size=12)

    # 15
    s = deck.slide("后续计划：先建立证据，再讨论部署",
                   "冻结当前对照基线；事件特征仅保留可选实验，不直接修改01默认模型输入",
                   "复现入口：09_方向优化实验.py；结果：OPT9_对比汇总.csv、OPT9_event_decision.json；本PPT由10_汇报PPT生成.py生成")
    table(s, ["优先级 / 阶段", "具体工作", "交付与通过门槛"], [
        ["P0：数据与口径", "发布时点/修订值审计；事件来源与标注版本冻结；同H10重训XGB",
         "同日历、同标签、同切分；独立逐日指标对账"],
        ["P1：概率与标签", "Platt/温度校准；方向与幅度分头；成熟OOS标签上开发期验证",
         "开发期超同子集多数类；配对分块区间为正；多数seed改善"],
        ["P2：稳健性复核", "同周期XGB＋NN融合；低置信弃权；按趋势/波动状态分层",
         "准确率与覆盖率同时报告；冻结方案后新时期盲测"],
        ["P3：影子运行", "固定信号时点；10日持有期；成本、换手、回撤与概率漂移监控",
         "成本后收益与风险合格；不以准确率替代交易可行性"],
    ], .66, 2.12, [2.4, 5.2, 4.4], .83, 13.5)
    box(s, .7, 6.47, 11.9, .43, NAVY)
    text(s, "决策建议：保持研究对照，不提前部署；优先验证校准与同周期公平对照。",
         .88, 6.54, 11.5, .31, 16, WHITE, True)

    assert len(deck.prs.slides) == 15
    DEST.mkdir(exist_ok=True)
    deck.prs.save(PPT)
    print(f"已生成：{PPT}")
    print("页数：15")
    for i, title in enumerate(deck.titles, 1):
        print(f"{i:02d} {title}")


if __name__ == "__main__":
    main()
